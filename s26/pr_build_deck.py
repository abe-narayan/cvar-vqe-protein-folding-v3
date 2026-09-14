"""PR lane, Sprint 26: build `vqe_research_overview.pptx` (repository root) from artefacts.

    python s26/pr_build_deck.py              # regenerate figures, build the deck, verify it
    python s26/pr_build_deck.py --no-figures # skip the figure step

What it reads: `s26/pr_values.py` (every number, each with its artefact path), `s26/pr_notes.md`
(the spoken text, with {TOKEN} placeholders), `s26/figures/*.png` (the A2 and A4 figures from
lane Q; the overlays, width sweep and ladder from `s26/pr_figures.py`).

What it writes: the deck; `s26/pr_values.json` (token -> value, path, basis, status);
`s26/pr_verify_dump.txt` (every text frame and notes frame of the reopened deck);
`s26/pr_verify.txt` (slide count, banned-word and dash grep, spoken word count per slide).

Style (campaign prompt Part 6): dark theme, background RGB(18, 18, 24), off-white text, one
accent colour, figures on a white plate, title + body + notes on every slide, one font.
No U+2014 or U+2013 anywhere; none of the six banned words. The basis of every RMSD is named.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)
from s26 import pr_values

DECK = os.path.join(ROOT, "vqe_research_overview.pptx")
FIG = os.path.join(ROOT, "s26", "figures")
NOTES_MD = os.path.join(ROOT, "s26", "pr_notes.md")

BG = RGBColor(18, 18, 24)
TEXT = RGBColor(236, 236, 240)
MUTED = RGBColor(168, 168, 182)
ACCENT = RGBColor(242, 169, 59)
PANEL = RGBColor(30, 30, 40)
FONT = "Calibri"
W, H = 13.333, 7.5
BANNED = ("genuinely", "honestly", "leverage", "robust", "delve", "underscore")


# --------------------------------------------------------------------------- notes and tokens
_TOKEN = re.compile(r"\{([A-Z][A-Z0-9_]*)(?:\[(\d+)\])?(?::([^}]*))?\}")


def fill(text, V, used):
    """Replace {TOKEN}, {TOKEN:spec}, {TOKEN[i]:spec} by the registry value."""
    def rep(m):
        tok, idx, spec = m.group(1), m.group(2), m.group(3)
        if tok not in V:
            raise KeyError(f"unknown token {tok}")
        used.add(tok)
        v = V[tok]["value"]
        if idx is not None:
            v = v[int(idx)]
        if isinstance(v, bool):
            return str(v)
        if isinstance(v, (int, float)) and spec:
            return format(v, spec)
        if isinstance(v, int):
            return f"{v:,}"
        if isinstance(v, float):
            return format(v, ".4f")
        if isinstance(v, list):
            if all(isinstance(x, int) and not isinstance(x, bool) for x in v):
                return ", ".join(str(x) for x in v)
            return "[" + ", ".join(format(float(x), spec or "+.4f") for x in v) + "]"
        return str(v)
    return _TOKEN.sub(rep, text)


def parse_notes():
    """{slide_no: {'title':..., 'spoken':..., 'also':...}} from pr_notes.md."""
    text = open(NOTES_MD, encoding="utf-8").read()
    out = {}
    for m in re.finditer(r"^## Slide (\d+) -- (.*?)$\n(.*?)(?=^## Slide |\Z)", text, re.M | re.S):
        no, title, body = int(m.group(1)), m.group(2).strip(), m.group(3)
        blocks = {}
        for b in re.finditer(r"^### (spoken|also)\s*$\n(.*?)(?=^### |\Z)", body, re.M | re.S):
            blocks[b.group(1)] = " ".join(line.strip() for line in b.group(2).strip().splitlines() if line.strip())
        out[no] = dict(title=title, spoken=blocks.get("spoken", ""), also=blocks.get("also", ""))
    return out


def notes_text(no, notes, V):
    """The full notes frame: spoken text, the 'also' block, then SOURCES for every token used."""
    used = set()
    spoken = fill(notes[no]["spoken"], V, used)
    also = fill(notes[no]["also"], V, used) if notes[no]["also"] else ""
    lines = ["SPOKEN:", spoken]
    if also:
        lines += ["", "ALSO (not spoken):", also]
    lines += ["", "SOURCES (not spoken; token = value <- artefact path [basis]):"]
    for tok in sorted(used):
        d = V[tok]
        val = d["value"]
        sval = format(val, ".6g") if isinstance(val, float) else str(val)
        basis = f" [{d['basis']}]" if d["basis"] else ""
        status = "" if d["status"] == "SOURCED" else f" ({d['status']})"
        lines.append(f"{tok} = {sval} <- {d['path']}{basis}{status}")
    return "\n".join(lines), spoken, sorted(used)


# --------------------------------------------------------------------------- drawing helpers
def new_slide(prs, title, no, footer=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bgf = slide.background.fill
    bgf.solid(); bgf.fore_color.rgb = BG
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(W - 1.0), Inches(0.95))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.name = FONT; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = ACCENT
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(1.22), Inches(W - 0.5), Inches(1.22))
    line.line.color.rgb = ACCENT; line.line.width = Pt(1.25)
    ft = slide.shapes.add_textbox(Inches(0.5), Inches(7.02), Inches(W - 1.0), Inches(0.35))
    ftf = ft.text_frame; p = ftf.paragraphs[0]
    r = p.add_run(); r.text = f"{no} / 11    {footer}"; r.font.name = FONT; r.font.size = Pt(10); r.font.color.rgb = MUTED
    return slide


def add_text(slide, left, top, width, height, items, size=15, color=TEXT, bullet=True, spacing=4, anchor=None):
    """items: list of str or (str, dict) with keys size/color/bold/bullet/indent."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    first = True
    for it in items:
        s, o = (it, {}) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(o.get("space", spacing))
        b = o.get("bullet", bullet)
        ind = o.get("indent", 0)
        txt = ("• " if b else "") + s
        if ind:
            txt = "    " * ind + txt
        r = p.add_run(); r.text = txt
        r.font.name = FONT; r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", False)
        r.font.color.rgb = o.get("color", color)
    return tb


def add_picture(slide, path, left, top, max_w, max_h, plate=True):
    """Place a PNG inside the box (left, top, max_w, max_h) keeping its aspect, on a white plate."""
    im = Image.open(path); w, h = im.size
    scale = min(max_w / w, max_h / h)
    pw, ph = w * scale, h * scale
    if plate:
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left - 0.05), Inches(top - 0.05), Inches(pw + 0.1), Inches(ph + 0.1))
        rect.fill.solid(); rect.fill.fore_color.rgb = RGBColor(255, 255, 255); rect.line.fill.background()
    slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(pw), height=Inches(ph))
    return pw, ph


def add_box(slide, left, top, width, height, text, size=11, fill_rgb=PANEL, line_rgb=ACCENT, color=TEXT, bold_first=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill_rgb
    shp.line.color.rgb = line_rgb; shp.line.width = Pt(1.0)
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line; r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold_first and i == 0
    return shp


def arrow(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = ACCENT; c.line.width = Pt(1.5)
    return c


def F(V, tok, spec=None):
    return pr_values.fmt(V, tok, spec)


# --------------------------------------------------------------------------- slides
def slide_01(prs, V):
    s = new_slide(prs, "A verified CVaR-VQE inside a peptide structure predictor", 1,
                  "every number on these slides has an artefact path in the speaker notes")
    add_text(s, 0.8, 1.9, 11.7, 1.0, [("What I built, what it measured, and where I want to go", dict(size=24, bullet=False, color=TEXT))])
    add_text(s, 0.8, 3.0, 11.7, 3.2, [
        (f"{F(V,'N_TARGETS')} held-out peptides, 9 to 16 residues, in 5 pinned folds; a sealed 60-target benchmark spent once", dict(size=17)),
        (f"Result: mean CA-RMSD {F(V,'ARM_MEAN','.4f')} A on the built chain (the production basis); "
         f"{F(V,'AVG_MEAN','.4f')} A is the point-cloud intermediate and is not a structure", dict(size=17)),
        (f"Quantum component: {F(V,'Q_QUBITS')} qubits, {F(V,'Q_LAYERS')} layers, {F(V,'Q_PARAMS')} parameters, exact statevector, "
         f"exact parameter-shift gradient; verified, then measured: no detectable accuracy contribution", dict(size=17)),
        ("What survives: the measuring instrument, and an exact chain from force-field locality to gradient variance", dict(size=17)),
        ("Presented to a postdoc in variational quantum algorithms. Sprint 26 deck, built from artefacts by s26/pr_build_deck.py", dict(size=13, color=MUTED, bullet=False)),
    ], size=17)
    return s


def slide_02(prs, V):
    s = new_slide(prs, "Why 9 to 16 residue peptides are hard", 2, "sources: s26/results/a_c26_phi_mae.json, s23/results/errdecomp.json, s24/results/priorladder.json, s12/results/s14_ladder.json")
    add_text(s, 0.6, 1.45, 12.1, 5.4, [
        ("Almost no sequence signal at this length", dict(bold=True, color=ACCENT, bullet=False)),
        (f"phi from the full sequence context: {F(V,'PHI_SEQ','.1f')} deg mean error over {F(V,'PHI_N_RES')} residues; "
         f"a sequence-blind corpus marginal: {F(V,'PHI_BLIND','.1f')} deg (ORACLE diagnostic, no RMSD)", {}),
        ("The retrieved candidates share their mistake", dict(bold=True, color=ACCENT, bullet=False)),
        (f"{F(V,'COMMON_MODE','.0%')} of the pool's error is common to every member (mean over {F(V,'N_TARGETS')} targets); "
         "averaging removes only the rest", {}),
        ("The distance prior is the only steep lever", dict(bold=True, color=ACCENT, bullet=False)),
        (f"interpolating the prior toward the true distances moves the endpoint {F(V,'PRIOR_SLOPE','.2f')} A per unit at the origin; "
         f"{F(V,'PRIOR_GAMMA_FOR_3A','.1%')} of the way reaches 3.0 A (ORACLE; point-cloud basis)", {}),
        ("Scale", dict(bold=True, color=ACCENT, bullet=False)),
        (f"a constant alpha-helix (zero information) scores {F(V,'HELIX','.3f')} A; the pipeline is {F(V,'HELIX_MINUS_ARM','.3f')} A better "
         f"({F(V,'ARM_MEAN','.4f')} A); both built chains", {}),
        ("The instrument", dict(bold=True, color=ACCENT, bullet=False)),
        (f"{F(V,'N_TARGETS')} cluster-disjoint targets, 5 pinned folds, paired per target; MDE = {F(V,'MDE_FACTOR','.4f')} x SE per comparison; "
         "fold-clustered CI beside every iid CI; pre-registered falsifiers; a corrections ledger", {}),
    ], size=15, spacing=3)
    return s


def slide_03(prs, V):
    s = new_slide(prs, "The pipeline, stage by stage", 3, "sources: bench_results/baseline_tuning126.json, s26/results/e_reproduce.json, s26/EXAMINATION.md section B")
    row1 = [("sequence\nESM-2 embedding", ""),
            ("distogram\n17-bin distance posterior\nleave-fold-out", ""),
            (f"BLOSUM62 retrieval\nK = 500 real windows\n{F(V,'CORPUS_PEPTIDES')} peptides + {F(V,'CORPUS_FRAGMENTS')} fragments", ""),
            ("Bayes-risk score\nkeep the top 75\n(top 128 when the VQE runs)", "")]
    row2 = [(f"CVaR-VQE selector\n{F(V,'Q_QUBITS')} qubits, {F(V,'Q_LAYERS')} layers, {F(V,'Q_PARAMS')} params\nOFF in production", "off"),
            (f"coordinate average\npoint cloud {F(V,'AVG_MEAN','.4f')} A\nnot a structure (bond {F(V,'BOND_AVG_MEAN','.2f')} A, worst {F(V,'BOND_AVG_WORST','.2f')})", ""),
            (f"ideal-geometry projection\nbuilt chain {F(V,'ARM_MEAN','.4f')} A\nTHE RESULT", "result"),
            (f"optional AMBER relaxation\nff14SB/GBn2, restrained\n{F(V,'FULL_MEAN','.4f')} A: a validity step", "")]
    bw, bh, gap, x0 = 2.75, 1.25, 0.32, 0.65
    for i, (t, kind) in enumerate(row1):
        x = x0 + i * (bw + gap)
        add_box(s, x, 1.65, bw, bh, t, size=11.5)
        if i < 3:
            arrow(s, x + bw, 1.65 + bh / 2, x + bw + gap, 1.65 + bh / 2)
    # the turn from row 1 to row 2
    arrow(s, x0 + 3 * (bw + gap) + bw / 2, 1.65 + bh, x0 + 3 * (bw + gap) + bw / 2, 3.35)
    arrow(s, x0 + 3 * (bw + gap) + bw / 2, 3.35, x0 + bw / 2, 3.35)
    arrow(s, x0 + bw / 2, 3.35, x0 + bw / 2, 3.6)
    for i, (t, kind) in enumerate(row2):
        x = x0 + i * (bw + gap)
        line_rgb = MUTED if kind == "off" else ACCENT
        fill_rgb = RGBColor(40, 40, 52) if kind == "result" else PANEL
        add_box(s, x, 3.6, bw, bh, t, size=11.5, line_rgb=line_rgb, fill_rgb=fill_rgb, bold_first=(kind == "result"))
        if i < 3:
            arrow(s, x + bw, 3.6 + bh / 2, x + bw + gap, 3.6 + bh / 2)
    add_text(s, 0.65, 5.15, 12.0, 1.8, [
        ("The output is an average of the kept set, so it tracks the set mean, not the best member; a selector that only removes "
         "members from a ranked list cannot move it much.", {}),
        (f"Hamiltonian of the selector: H = diag(zrank(scores)) on the top 128; the same rank ladder on every target to "
         f"{F(V,'Q_LADDER_WORST_FRAC','.2%')} of its range. The production run has quantum = False; the {F(V,'ARM_MEAN','.4f')} A never passes through the VQE.", {}),
        (f"Reproduction: a fresh run of the instrument over all {F(V,'N_TARGETS')} stored records disagrees with the record by "
         f"{F(V,'REPRO_MAX_DISAGREE','.1f')} on every basis; 1-worker and 8-worker runs agree to {F(V,'SCIENCE_DELTA_MAX','.0e')}.", {}),
    ], size=13.5, spacing=3)
    return s


def slide_04(prs, V):
    s = new_slide(prs, f"Results on {F(V,'N_TARGETS')} held-out targets (built chain), with two overlays", 4,
                  "overlays: bench_results/cache/1fc9f2dcf489e2fb/<pdb>.json 'ca' vs s8/generate_univ/<pdb>.npz 'nat_ca', ORACLE-superposed for the figure")
    add_picture(s, os.path.join(FIG, "pr_overlay_1S9Z.png"), 0.55, 1.45, 6.1, 2.15)
    add_picture(s, os.path.join(FIG, "pr_overlay_9KAR.png"), 0.55, 3.75, 6.1, 3.15)
    add_text(s, 6.95, 1.4, 6.0, 5.6, [
        ("Built chain, production cache", dict(bold=True, color=ACCENT, bullet=False)),
        (f"mean {F(V,'ARM_MEAN','.4f')} A, median {F(V,'ARM_MEDIAN','.3f')}, best {F(V,'ARM_MIN','.3f')} ({F(V,'ARM_BEST_PDB')}), "
         f"worst {F(V,'ARM_MAX','.3f')} ({F(V,'ARM_WORST_PDB')})", {}),
        (f"{F(V,'ARM_FRAC2','.1%')} of targets under 2 A, {F(V,'ARM_FRAC3','.1%')} under 3 A", {}),
        (f"point cloud {F(V,'AVG_MEAN','.4f')} A (intermediate); relaxed {F(V,'FULL_MEAN','.4f')} A", {}),
        ("Controls and oracles on the same targets", dict(bold=True, color=ACCENT, bullet=False)),
        (f"random 75-subset {F(V,'RANDOM75','.4f')} A on the point-cloud basis (production point cloud {F(V,'AVG_MEAN','.4f')})", {}),
        (f"constant alpha-helix {F(V,'HELIX','.3f')} A (built chain)", {}),
        (f"ORACLE best single window: pool {F(V,'POOL_BEST','.3f')} A, shipped top-75 {F(V,'TOPM_BEST','.3f')} A", {}),
        (f"The hard target {F(V,'HARD_PDB')}: pool best {F(V,'HARD_POOL_BEST','.2f')} A, top-75 best {F(V,'HARD_TOPM_BEST','.2f')} A (ORACLE): "
         "the score filter, not the search, discards the good window", {}),
        ("Sealed benchmark, 60 targets, spent once", dict(bold=True, color=ACCENT, bullet=False)),
        (f"full system minus shipped baseline: {F(V,'BENCH_DELTA','+.4f')} A paired, CI [{F(V,'BENCH_CI_LO','+.3f')}, {F(V,'BENCH_CI_HI','+.3f')}], "
         f"{F(V,'BENCH_WL')}. No validated improvement.", {}),
    ], size=13, spacing=2)
    return s


def slide_05(prs, V):
    s = new_slide(prs, "The quantum component: verified, then measured", 5, "sources: s25/results/q_verify.json, q_gibbs.json, q_alpha.json; s25/QUANTUM.md")
    add_text(s, 0.6, 1.4, 6.05, 5.6, [
        ("Verified", dict(bold=True, color=ACCENT, bullet=False, size=16)),
        (f"{F(V,'Q_QUBITS')} qubits, {F(V,'Q_LAYERS')} layers of RY + CNOT chain + ring, {F(V,'Q_PARAMS')} parameters, "
         f"{F(V,'Q_DIM')} candidate states; exact dense statevector", {}),
        (f"vs an independent dense simulator: max |dp| = {F(V,'Q_SV_ERR')}", {}),
        (f"exact parameter-shift gradient vs finite differences: cos {F(V,'Q_PS_COS')}, rel. err. {F(V,'Q_PS_REL')} "
         f"(free-energy gradient {F(V,'Q_FE_REL')})", {}),
        (f"sampling in the trained path: {F(V,'Q_SAMPLING')}; RNG: {F(V,'Q_RNG')}", {}),
        ("objective F = CVaR_alpha(E) - T H(p); H = diag(zrank(scores)): the rank ladder, target-independent to "
         f"{F(V,'Q_LADDER_WORST_FRAC','.2%')} of its range ({F(V,'Q_LADDER_N_CHECKED')} targets checked)", {}),
        (f"theorem: the CVaR tail is a subset of a prefix of the energy order; {F(V,'Q_CELLS')} cells, {F(V,'Q_VIOLATIONS')} violations, "
         f"exact equality on {F(V,'Q_EQUALITY')} full-support cells", {}),
    ], size=13, spacing=3)
    add_text(s, 6.95, 1.4, 5.9, 5.6, [
        ("Measured", dict(bold=True, color=ACCENT, bullet=False, size=16)),
        (f"the optimiser trains: beats best-of-{F(V,'Q_BEST_OF_N')} untrained draws at every T; closes "
         f"{F(V,'Q_GAP_MIN','.0%')} to {F(V,'Q_GAP_MAX','.0%')} of the free-energy gap", {}),
        (f"it stops {F(V,'Q_KL_NATS','.3f')} nats (TV {F(V,'Q_TV','.3f')}) from its Gibbs optimum, broader than optimal "
         f"({F(V,'Q_H_TRAINED','.2f')} vs {F(V,'Q_H_GIBBS','.2f')} bits); the endpoint cannot tell: {F(V,'Q_CIRC_VS_GIBBS_X','.2f')} x MDE", {}),
        (f"selector vs the plain argmin: {F(V,'Q_VS_ARGMIN','+.4f')} A at {F(V,'Q_VS_ARGMIN_X','.2f')} x MDE, "
         f"{F(V,'Q_VS_ARGMIN_W')}W/{F(V,'Q_VS_ARGMIN_L')}L/{F(V,'Q_VS_ARGMIN_T')}T (selection basis): {F(V,'Q_VS_ARGMIN_VERDICT')}", {}),
        (f"vs an exact Boltzmann weighting at the same T: {F(V,'Q_VS_BOLTZ','+.4f')} A", {}),
        (f"alpha = 1 on folds {F(V,'Q_ALPHA1_FOLDS')}: {F(V,'Q_SHARE_NO_TAIL','.1%')} of targets carry no tail constraint", {}),
        ("off in production; the theorem says it can delete a top-m member, never add one", {}),
        ("No quantum advantage is claimed. The register is 7 qubits and every state is simulated exactly.", dict(bold=True)),
    ], size=13, spacing=3)
    return s


def slide_06(prs, V):
    s = new_slide(prs, "Gradient variance at depth 3, and the algebra (not a barren-plateau claim)", 6,
                  "sources: s25/results/q_plateau.json (left), s26/results/q_dla.json (right, lane Q's A2 figure), s13/results/geo_kernel.json")
    add_picture(s, os.path.join(FIG, "pr_width_sweep.png"), 0.5, 1.4, 7.0, 3.05)
    add_picture(s, os.path.join(FIG, "a2_dla_dimension.png"), 7.75, 1.4, 5.1, 3.3)
    add_text(s, 0.55, 4.65, 12.3, 2.35, [
        (f"Width sweep, n = {F(V,'SWEEP_N_MIN')}..{F(V,'SWEEP_N_MAX')}, depth 3, exact gradients, theta ~ N(0, 0.6^2): fitted log2 Var per qubit "
         f"{F(V,'SLOPE_LIN','.3f')} (linear cost), {F(V,'SLOPE_A025_T0','.3f')} / {F(V,'SLOPE_A01_T0','.3f')} (CVaR alpha 0.25 / 0.10), "
         f"{F(V,'SLOPE_A1_T03','.3f')} / {F(V,'SLOPE_A025_T03','.3f')} (deployed, T = 0.3). Reference: the depth-8 decay base "
         f"{F(V,'GEO_BASE_DEPTH8','.3f')} per qubit ({F(V,'GEO_SLOPE_DEPTH8','.2f')} in log2).", {}),
        (f"What it says: at this depth the decay is slower than the 2-design rate, and the CVaR non-linearity flattens it "
         f"(variance ratio CVaR/linear {F(V,'RATIO_N7_A025','.2f')} at n = 7, {F(V,'RATIO_N13_A025','.2f')} at n = 13). "
         f"Depth at n = 7 saturates from L = 4 ({F(V,'DEPTH_L3','.1e')} at L = 3, {F(V,'DEPTH_L12','.1e')} at L = 12).", {}),
        (f"What it does not say: it is not a 2-design result and not 'no barren plateau'. {F(V,'Q_PARAMS')} parameters against dim so(128) = "
         f"{F(V,'DLA_SO128')}; the dynamical Lie algebra is the full so(2^n) from depth 2 at n = 7 (depth 1: abelian, dim {F(V,'DLA_N7_L1')}), "
         f"so nothing algebraic protects the ansatz at scale. Draws {F(V,'SWEEP_DRAWS_MAX')} to {F(V,'SWEEP_DRAWS_MIN')} per width: relative SE 9 to 16%.", {}),
    ], size=12.5, spacing=3)
    return s


def slide_07(prs, V):
    s = new_slide(prs, "Where the remaining accuracy lives", 7,
                  "sources: s25/results/phys_suite.json (point cloud), s24/results/priorladder.json, s23/results/errdecomp.json, s26/results/ph_c3_stage1.json")
    add_picture(s, os.path.join(FIG, "pr_ladder.png"), 0.5, 1.4, 7.0, 4.1)
    add_text(s, 7.75, 1.35, 5.2, 5.7, [
        ("Closed by measurement", dict(bold=True, color=ACCENT, bullet=False)),
        (f"selection and ranking inside the pool: the pool holds a {F(V,'POOL_BEST','.2f')} A window and the top-75 a "
         f"{F(V,'TOPM_BEST','.2f')} A one (ORACLE); no native-free ranker finds them", {}),
        (f"physics as a selector: Legacy {F(V,'LEGACY_VS_RANDOM','+.3f')} A and AMBER {F(V,'AMBER_VS_RANDOM','+.3f')} A WORSE than a random "
         f"75-subset, {F(V,'LEGACY_VS_RANDOM_FOLDS')}/5 folds each (point-cloud basis; fold CIs in the notes)", {}),
        (f"physics as a relaxer: AMBER costs {F(V,'C3_AMBER_VS_NONE','+.4f')} A against the built chain ({F(V,'C3_AMBER_VS_NONE_X','.2f')} x MDE) "
         f"and is {F(V,'C3_AMBER_VS_RANDOM','+.4f')} A worse than a random move of its own size (fold CI {F(V,'C3_AMBER_VS_RANDOM_CI','+.4f')}, "
         "Type-M zone: sign measured, size an upper bound); a validity step on 124 of 126", {}),
        (f"physics as a steric reject filter (1e4 kcal/mol, point cloud): {F(V,'REJECT_R_VS_ANCHOR','+.3f')} A worse than the shipped top-75, "
         f"median {F(V,'REJECT_R_VS_ANCHOR_MEDIAN','+.3f')} (tail-carried; Type-M)", {}),
        ("search: the exact argmin over the whole latent ties the pool (S21)", {}),
        ("Open", dict(bold=True, color=ACCENT, bullet=False)),
        (f"the prior: {F(V,'PRIOR_SLOPE','.2f')} A per unit toward the true distances; a perfect prior through the same pipeline reaches "
         f"{F(V,'PRIOR_PERFECT','.3f')} A from {F(V,'PRIOR_GAMMA0','.3f')} (ORACLE, point cloud)", {}),
        (f"the pool's error is {F(V,'COMMON_MODE','.0%')} common-mode: nothing downstream of retrieval can remove most of it", {}),
        ("If accuracy moves, it moves through a better distance prior.", dict(bold=True)),
    ], size=12.5, spacing=3)
    add_text(s, 0.5, 5.6, 7.0, 1.4, [
        (f"Basis on every bar. Built chain: production {F(V,'ARM_MEAN','.4f')}, helix {F(V,'HELIX','.3f')}, torsion predictor {F(V,'TORS','.3f')}. "
         f"Point cloud: production {F(V,'AVG_MEAN','.4f')}, random-75 {F(V,'RANDOM75','.4f')}, perfect prior {F(V,'PRIOR_PERFECT','.3f')}. "
         f"Single window: pool best {F(V,'POOL_BEST','.3f')}, top-75 best {F(V,'TOPM_BEST','.3f')}, whole library {F(V,'UNIVERSE_BEST','.3f')}.", dict(bullet=False)),
    ], size=10.5, color=MUTED)
    return s


def slide_08_pending(prs, V):
    s = new_slide(prs, "Direction A: ADAPT-VQE on this Hamiltonian (PENDING)", 8,
                  "PENDING: wording and verdict from s26/PROPOSAL_A.md when lane Q posts it; the facts shown are measured (A2 L27, A4 L35; Adversary L45, L47)")
    add_picture(s, os.path.join(FIG, "a4_variance_slopes.png"), 0.5, 1.4, 7.2, 3.0)
    add_text(s, 0.55, 4.5, 7.1, 2.5, [
        (f"A4 (lane Q's figure, s26/results/q_var.json): grown circuits at alpha = 1 hold {F(V,'A4_A1_DISTINCT_MIN')} to "
         f"{F(V,'A4_A1_DISTINCT_MAX')} distinct operators (product circuits, per the adapt sets of q_dla.json); their variance does not decay "
         f"with n (slopes {F(V,'A4_V_LIN','+.3f')} / {F(V,'A4_L2_LIN','+.3f')} at T = 0, {F(V,'A4_V_A1_T03','+.3f')} / {F(V,'A4_L2_A1_T03','+.3f')} at T = 0.3) "
         f"and is {F(V,'A4_RATIO_MIN','.1f')} to {F(V,'A4_RATIO_MAX','.0f')} x the fixed ansatz's at matched n. A large gradient from a product circuit is not trainability.", {}),
        (f"the one non-trivial grown family (alpha 0.25, pool L2, {F(V,'A4_L2_A025_DISTINCT_MIN')} to {F(V,'A4_L2_A025_DISTINCT_MAX')} distinct 2-local strings) "
         f"decays at {F(V,'A4_L2_A025_T03','.3f')} per qubit vs the fixed {F(V,'A4_FIXED_A025_T03','.3f')}; both far from a 2-design's -1; no CI on the difference is on disk (L47)", {}),
    ], size=11, spacing=3)
    add_text(s, 7.95, 1.4, 4.95, 5.6, [
        ("What the proposal says", dict(bold=True, color=ACCENT, bullet=False)),
        ("replace the fixed RY/CNOT ansatz with an adaptively grown one (qubit-ADAPT pools V and L2)", {}),
        ("A2, final: the algebra offers nothing to grow into", dict(bold=True, color=ACCENT, bullet=False)),
        (f"dim(DLA) at n = 7: {F(V,'DLA_N7_L1')} at depth 1 (abelian), {F(V,'DLA_N7_L2')} from depth 2 = dim so(128) = {F(V,'DLA_SO128')}; "
         f"pool V generates {F(V,'DLA_POOL_V_N7')}, pool L2 {F(V,'DLA_POOL_L2_N7')}; ADAPT-selected sets: dim {F(V,'ADAPT_A1_DIM')} at alpha = 1, "
         f"{F(V,'ADAPT_A025_L2_DIM')} at alpha = 0.25 (L2)", {}),
        ("The target is a product state", dict(bold=True, color=ACCENT, bullet=False)),
        (f"KL(Gibbs || product of its marginals) = {F(V,'PROD_KL_IDEAL','.1e')} on the ideal ladder, {F(V,'PROD_KL_1A13','.1e')} on 1A13; "
         f"a 7-parameter RY layer reaches the Gibbs state to {F(V,'PROD_KL_RY7','.1e')} nats, the deployed 21-parameter circuit stops at {F(V,'PROD_KL_FIXED','.3f')}", {}),
        ("A1, endpoint: running (fixed vs ADAPT-grown, built chain primary); expected null by the set-equality theorem", {}),
        ("Verdict: PENDING s26/PROPOSAL_A.md. Nothing here is a width-scaling argument for ADAPT on this Hamiltonian (L35).", dict(bold=True)),
    ], size=11.5, spacing=3)
    return s


def slide_09(prs, V):
    s = new_slide(prs, "Direction B: a learned folding model as the prior, and its replacement", 9,
                  "from s26/PROPOSAL_B.md (lane P, DRAFT; B1 final) and s26/PROPOSAL_B_REPLACEMENT.md (lane Q); the A2 and A4 figures are on slides 6 and 8")
    add_text(s, 0.6, 1.4, 6.1, 5.6, [
        ("What the proposal says", dict(bold=True, color=ACCENT, bullet=False)),
        ("Replace or augment the shipped distance prior with the output of a large pretrained folding model, because the prior's "
         "accuracy is the only steep lever the record measured.", {}),
        ("B1, final: the folding model cannot run on this machine", dict(bold=True, color=ACCENT, bullet=False)),
        (f"no checkpoint on disk; {F(V,'B1_DOWNLOAD_GB','.2f')} GB to download", {}),
        ("openfold and omegaconf absent; openfold needs nvcc and Python <= 3.9; this box is Python 3.13, CPU-only torch", {}),
        (f"{F(V,'B1_RESIDENT_GB','.2f')} GB resident as fair-esm loads it (fp16 LM, fp32 trunk) against {F(V,'B1_HEADROOM_GB','.1f')} GB of headroom; "
         f"fp16 everywhere is still {F(V,'B1_FP16_GB','.1f')} GB", {}),
        ("B2, pending: the feasible-scale ladder", dict(bold=True, color=ACCENT, bullet=False)),
        ("the ESM-2 contact head alone as a prior input; the 8M language model against the 650M (rungs conly and esm8m of the C2 ladder)", {}),
        ("B3, pending: where the pipeline beats sequence-only, and is that set recognisable native-free?", dict(bold=True, color=ACCENT, bullet=False)),
        (f"built chains: pipeline {F(V,'ARM_MEAN','.4f')}, sequence-only torsion predictor {F(V,'TORS','.4f')}, constant helix {F(V,'HELIX','.4f')} A; "
         f"torsion minus pipeline {F(V,'B3_TORS_MINUS_ARM','+.4f')} (SE {F(V,'B3_TORS_MINUS_ARM_SE','.4f')}, MDE {F(V,'B3_TORS_MINUS_ARM_MDE','.4f')}, "
         f"{F(V,'B3_TORS_W')}W/{F(V,'B3_TORS_L')}L); the classifier test is pending", {}),
        ("Verdict: PENDING B2/B3. If null: REPLACE with the trainability paper (right).", dict(bold=True)),
    ], size=11.5, spacing=3)
    add_text(s, 6.95, 1.4, 5.95, 5.6, [
        ("The replacement: publish the trainability result", dict(bold=True, color=ACCENT, bullet=False)),
        ("chain geometry fixes which residues a distance term can depend on (exact theorem); that fixes the energy's Pauli spectrum "
         "once raw AMBER's clash spike is conditioned away", {}),
        (f"the spectrum times the circuit's own kernel predicts the measured gradient variance with no free parameter: median "
         f"measured/predicted {F(V,'PAULI_RATIO_LEGACY','.3f')} over {F(V,'PAULI_N_CELLS')} cells", {}),
        ("at depth >= 3 the kernel is flat in Pauli weight: the force field's locality is not what limits training; the width is", {}),
        (f"the Lie algebra is the full so(2^n) from depth 2 (dim {F(V,'DLA_SO128')} at n = 7): nothing algebraic protects the ansatz at scale; "
         f"n = {F(V,'SWEEP_N_MIN')}..{F(V,'SWEEP_N_MAX')} at depth 3 is the shallow regime", {}),
        (f"the deployed Gibbs target is a product state (KL to its product of marginals {F(V,'PROD_KL_1A13','.1e')} on 1A13); an ADAPT-grown "
         "circuit on it grows no entanglement, so its gradients do not decay because there is nothing to train", {}),
        (f"the optimiser trains ({F(V,'Q_GAP_MIN','.0%')} to {F(V,'Q_GAP_MAX','.0%')} of the gap) and the readout cannot tell "
         f"({F(V,'Q_CIRC_VS_GIBBS_X','.2f')} x MDE); the tail is a subset of the classical prefix ({F(V,'Q_CELLS')} cells, {F(V,'Q_VIOLATIONS')} violations)", {}),
        ("Two claims the paper will not make", dict(bold=True, color=ACCENT, bullet=False)),
        ("a quantum advantage (7 qubits, exact simulation, selection classical by theorem, a product-state target)", {}),
        ("a barren plateau from a small gradient; 'no plateau' is said only as 'at n <= 13, depth 3, this ansatz'", {}),
        ("Venue: noiseless exact simulation throughout; a submission adds a noise model, width beyond 13, a 2-design control, error bars", dict(color=MUTED)),
    ], size=11.5, spacing=3)
    return s


def slide_10(prs, V):
    s = new_slide(prs, "Direction C: learn a better distance prior; physics for validity, not accuracy", 10,
                  "from s26/PROPOSAL_C.md (DRAFT, lane P) and s26/C3_RESULT.md (lane PH, L39, Adversary L46); verdict PENDING C2 to C5")
    add_text(s, 0.6, 1.4, 6.2, 5.6, [
        ("What the proposal says", dict(bold=True, color=ACCENT, bullet=False)),
        ("The prior bounds the accuracy; improve it by learning: better inputs (C2), physics kept honest (C3), routers for the "
         "per-target set size (C4), a learned correction of the common-mode error (C5); C1 reproduces the closures first.", {}),
        ("Not 'learn a better ranker': ranking inside the pool is closed at four levels; the prior is the only steep lever.", {}),
        ("C1, final: the closures reproduce to the third decimal", dict(bold=True, color=ACCENT, bullet=False)),
        (f"set-transformer ranker, real features: {F(V,'C1_S12_REAL_N8','.3f')} (n = 8) to {F(V,'C1_S12_REAL_FULL','.3f')} A (full); "
         f"with a leaked label {F(V,'C1_S12_LEAK_N8','.3f')} to {F(V,'C1_S12_LEAK_FULL','.3f')} (ORACLE; weighted-average basis)", {}),
        (f"perfect ranker inside the shipped top-25: {F(V,'C1_S17_BAND_BEST','.3f')} A vs random member {F(V,'C1_S17_RAND','.3f')} "
         f"and distogram argmin {F(V,'C1_S17_DIST_PICK','.3f')} (single window)", {}),
        ("C2, running: the prior-input ladder", dict(bold=True, color=ACCENT, bullet=False)),
        (f"{F(V,'C2_N_RUNGS')} rungs (no ESM, contact head only, pca32, pca32f, pca128, raw, 8M vs 650M, wide, pairnet, mix); "
         f"{F(V,'C2_RUNGS_TRAINED')} rungs have all five fold models on disk; evaluations running", {}),
        ("falsifier, fixed before the first run: beat the shipped posterior on the built chain by more than its own MDE, "
         "fold-clustered CI excluding zero, 5/5 folds", {}),
        (f"anchor: selection and point cloud reproduce the production cache ({F(V,'C2_ANCHOR_SEL_MAXABS','.0e')}, "
         f"{F(V,'C2_ANCHOR_CLOUD_MAXABS','.1e')}); built chain on the rebuild basis {F(V,'C2_ANCHOR_ARM','.4f')} A (L57)", {}),
    ], size=12, spacing=3)
    add_text(s, 7.05, 1.4, 5.85, 5.6, [
        ("C3, final: refine-with-physics is a validity step, not an accuracy step", dict(bold=True, color=ACCENT, bullet=False)),
        (f"the production relaxation moves the built chain {F(V,'C3_MAG','.3f')} A and costs {F(V,'C3_AMBER_VS_NONE','+.4f')} A "
         f"(fold CI {F(V,'C3_AMBER_VS_NONE_CI','+.4f')}), {F(V,'C3_AMBER_VS_NONE_X','.2f')} x MDE (relaxed minus built chain)", {}),
        (f"a random move of the same size costs {F(V,'C3_RANDOM_VS_NONE','+.4f')} A {F(V,'C3_RANDOM_VS_NONE_CI','+.4f')}; "
         f"AMBER minus random {F(V,'C3_AMBER_VS_RANDOM','+.4f')} {F(V,'C3_AMBER_VS_RANDOM_CI','+.4f')}, "
         f"{F(V,'C3_AMBER_VS_RANDOM_FOLDS')}/5 folds, {F(V,'C3_AMBER_VS_RANDOM_X','.2f')} x MDE: Type-M zone, sign measured, size an upper bound", {}),
        (f"AMBER minus a move of the same size toward a random pool member: {F(V,'C3_AMBER_VS_MEMBER','+.4f')} A "
         f"{F(V,'C3_AMBER_VS_MEMBER_CI','+.4f')}; ORACLE cosine of the displacement with the true residual: {F(V,'C3_COS','+.3f')}", {}),
        (f"what validity buys: {F(V,'C3_FRAC_E0_OVER_1E4','.1%')} of built chains start above 1e4 kcal/mol; {F(V,'C3_N_CONVERGED')} of 126 "
         f"converge below 1000 (9KAR ends at {F(V,'C3_E1_MAX','.0f')}); a validity step on 124 of 126, on 2BP4 and 9KAR it breaks a virtual bond", {}),
        ("C4, C5: pre-registered, pending", dict(bold=True, color=ACCENT, bullet=False)),
        ("routers on a feature set no previous router used (C4); predict and subtract the common mode (C5)", {}),
        ("Verdict: PENDING the ladder (KEEP AS STATED / KEEP WITH EDITS / REPLACE, per lane P)", dict(bold=True)),
    ], size=12, spacing=3)
    return s


def slide_pending(prs, V, no, title, waiting_for):
    s = new_slide(prs, title, no, f"PENDING: {waiting_for}")
    add_text(s, 0.8, 2.2, 11.7, 3.0, [
        (f"PENDING. This slide is filled from {waiting_for} when it lands.", dict(size=20, bullet=False, color=ACCENT)),
        ("Until then nothing on it is claimed.", dict(size=16, bullet=False)),
    ])
    return s


def slide_11(prs, V):
    s = new_slide(prs, "Goal and ask", 11, "the direction line in the notes is DRAFT until the coordinator's verdict entry in s26/LEDGER.md")
    add_text(s, 0.7, 1.45, 12.0, 5.5, [
        ("Goal", dict(bold=True, color=ACCENT, bullet=False, size=18)),
        ("Find out whether a better distance prior is obtainable from inputs this machine can compute, on the instrument I have, "
         "pre-registered, with fold-clustered intervals.", dict(size=15)),
        ("Publish the trainability chain: locality theorem, Pauli spectrum, kernel, width sweep, the Lie algebra and the product-circuit result.", dict(size=15)),
        ("Ask", dict(bold=True, color=ACCENT, bullet=False, size=18)),
        ("An independent study with those two threads, one measurement each.", dict(size=15)),
        ("One question: what does an adaptive ansatz do when the target Hamiltonian is not diagonal, and does the "
         "spectrum-times-kernel prediction extend to that case?", dict(size=15)),
        ("What I will not claim", dict(bold=True, color=ACCENT, bullet=False, size=18)),
        ("A quantum advantage (7 qubits, exact simulation, selection classical by theorem).", dict(size=15)),
        ("A barren plateau from a small gradient (a plateau is a decay with width; mine is a depth-3 measurement at n <= 13).", dict(size=15)),
        (f"An effect below its own minimum detectable size (MDE = {F(V,'MDE_FACTOR','.4f')} x SE, per comparison).", dict(size=15)),
    ], spacing=4)
    return s


# --------------------------------------------------------------------------- verification
def verify(V, spoken_by_slide):
    prs = Presentation(DECK)
    lines = []
    n = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"===== SLIDE {i} =====")
        for shp in slide.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                lines.append("[shape] " + shp.text_frame.text.replace("\n", " | "))
            if shp.shape_type == 13:
                lines.append(f"[picture] {shp.name} {shp.width.inches:.2f}x{shp.height.inches:.2f} in")
        lines.append("[notes] " + (slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""))
    dump = "\n".join(lines)
    with open(os.path.join(ROOT, "s26", "pr_verify_dump.txt"), "w", encoding="utf-8") as fh:
        fh.write(dump)
    checks = [f"deck: {os.path.relpath(DECK, ROOT)}", f"slide count: {n} (expected 11)"]
    em = dump.count("—"); en = dump.count("–")
    checks.append(f"U+2014 em dashes: {em}; U+2013 en dashes: {en}")
    low = dump.lower()
    for w in BANNED:
        checks.append(f"banned word '{w}': {len(re.findall(w, low))} occurrences")
    non_ascii = sorted({c for c in dump if ord(c) > 127})
    checks.append("non-ASCII characters present: " + " ".join(f"U+{ord(c):04X}" for c in non_ascii))
    for i, slide in enumerate(prs.slides, 1):
        has_title = any(shp.has_text_frame and shp.text_frame.text.strip() for shp in slide.shapes)
        has_notes = slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip() != ""
        wc = len(spoken_by_slide.get(i, "").split())
        flag = ""
        if i in (8, 9, 10):
            flag = "  (proposal slide: limit 250)" + ("  OK" if wc <= 250 else "  OVER LIMIT")
        checks.append(f"slide {i:2d}: text {'yes' if has_title else 'NO'}, notes {'yes' if has_notes else 'NO'}, spoken words {wc}{flag}")
    ok = (n == 11 and em == 0 and en == 0 and all(len(re.findall(w, low)) == 0 for w in BANNED)
          and all(len(spoken_by_slide.get(i, "").split()) <= 250 for i in (8, 9, 10)))
    checks.append("VERDICT: " + ("PASS" if ok else "FAIL"))
    out = "\n".join(checks)
    with open(os.path.join(ROOT, "s26", "pr_verify.txt"), "w", encoding="utf-8") as fh:
        fh.write(out + "\n")
    print(out)
    return ok


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-figures", action="store_true")
    args = ap.parse_args()
    if not args.no_figures:
        from s26 import pr_figures
        pr_figures.main()
    V = pr_values.load_values()
    with open(os.path.join(ROOT, "s26", "pr_values.json"), "w", encoding="utf-8") as fh:
        json.dump({k: dict(value=d["value"], path=d["path"], status=d["status"], basis=d["basis"], note=d["note"]) for k, d in V.items()},
                  fh, indent=1, default=str)
    notes = parse_notes()

    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    builders = {1: slide_01, 2: slide_02, 3: slide_03, 4: slide_04, 5: slide_05, 6: slide_06, 7: slide_07, 9: slide_09, 10: slide_10, 11: slide_11}
    if not os.path.exists(os.path.join(ROOT, 's26', 'PROPOSAL_A.md')):
        builders[8] = slide_08_pending   # the 'until then' state of step (d)
    pending = {8: ("Direction A: ADAPT-VQE on this Hamiltonian", "s26/PROPOSAL_A.md")}
    spoken_by_slide = {}
    manifest = {}
    for no in range(1, 12):
        if no in builders:
            s = builders[no](prs, V)
        else:
            title, waiting = pending[no]
            s = slide_pending(prs, V, no, title, waiting)
        text, spoken, used = notes_text(no, notes, V)
        s.notes_slide.notes_text_frame.text = text
        spoken_by_slide[no] = spoken
        manifest[no] = dict(title=notes[no]["title"], tokens=used, spoken_words=len(spoken.split()))
    prs.save(DECK)
    with open(os.path.join(ROOT, "s26", "pr_manifest.json"), "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, indent=1)
    print(f"wrote {os.path.relpath(DECK, ROOT)} ({len(prs.slides)} slides)")
    ok = verify(V, spoken_by_slide)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
